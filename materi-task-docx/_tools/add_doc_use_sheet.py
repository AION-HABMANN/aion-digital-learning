"""
Adds (or replaces) the sheet "Document use Day 1-8" in the One Stop Digital Learning Strukturplan xlsx.
It mirrors the plan's grid (UE rows, one column per day) and says, for each unit of each day, which document is used, which part of it,
and in which format (Presentation, Handout, Trainer only). Built from the same data as the lesson plan inside the Word documents.
Usage: python _tools/add_doc_use_sheet.py <output.xlsx>      (reads the plan xlsx next to the aion-digital-learning folder)
"""
import math
import os
import re
import sys
from copy import copy

import openpyxl
from openpyxl.styles import Alignment

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
import ablauf  # noqa: E402
import release  # noqa: E402

PLAN = os.path.join(os.path.dirname(ROOT), "One Stop Digital Learning - Strukturplan 11_F_128 (04-26).xlsx")
SHEET = "Document use Day 1-8"
P = "Habmann-Digital-Learning"
FILES = {
    1: ("UX-vs-UI-and-User-Centred-Design", "LearnFast-UX-Analysis-and-UX-Strategy-Memo", "SkillUp-and-LearnFast-Coaching-Feedback-and-Model-Solutions"),
    2: ("Successful-Platforms-Prototyping-Testing-and-Adaptive-Learning", "LearnPro-Prototype-and-Test-Plan-and-Innovation-Strategy", "LearnPro-Coaching-Feedback-and-Model-Solutions"),
    3: ("Learning-Psychology-and-Cognitive-Load", "EduCore-Cognitive-Load-Analysis-and-Learning-Experience-Strategy", "EduCore-Coaching-Feedback-and-Model-Solutions"),
    4: ("Motivation-and-Engagement-by-Design", "MotivaLearn-Engagement-Analysis-and-Engagement-Strategy", "MotivaLearn-Coaching-Feedback-and-Model-Solutions"),
    5: ("User-Centred-UX-Design", "LearnBase-User-Centred-Analysis-and-UX-Strategy", "LearnBase-Coaching-Feedback-and-Model-Solutions"),
    6: ("Personas-and-User-Journeys", "EduPath-Personas-Journeys-and-Segmentation-Strategy", "EduPath-Coaching-Feedback-and-Model-Solutions"),
    7: ("Accessibility-and-Inclusive-Design", "InclusiveLearn-Barrier-Analysis-and-Inclusive-UX-Strategy", "InclusiveLearn-Coaching-Feedback-and-Model-Solutions"),
    8: ("UX-UI-Structure-Visual-Hierarchy-and-Interaction", "StructLearn-Structure-Analysis-and-Basic-UX-UI-Structure", "StructLearn-Coaching-Feedback-and-Model-Solutions"),
}


def ids(cards):
    return ", ".join(c["id"] for c in cards)


def mins(cards):
    return sum(c["min"] for c in cards)


def blk(blocks, bid):
    b = blocks.get(bid)
    return f"Block {bid} ({b['kind']}, about {b['min']} min)" if b else f"Block {bid}"


def texts(day):
    cards, blocks = ablauf.parse_day(day)
    a = [c for c in cards if c["id"].startswith("A")]
    b = [c for c in cards if c["id"].startswith("B")]
    a_core = [c for c in a if c["kind"] == "Core"]
    a_opt = [c for c in a if c["kind"] == "Optional"]
    b_core = [c for c in b if c["kind"] == "Core"]
    case, role = ablauf.CASES[day], ablauf.ROLES[day]
    t = {}
    t[1] = f"MATERI A (Route 1)\nCore cards {ids(a_core)} (about {mins(a_core)} min)." + (f"\nOptional {ids(a_opt)}: handout only." if a_opt else "") + "\nFormat: Presentation + Handout"
    t[2] = f"TASK 1 · Part 1\n{blk(blocks, '1.1')}.\nFormat: Handout (or on the playground)"
    t[3] = (f"TASK 1 · Part 2\n{blk(blocks, '2.2')}: rating part; finished in UE5–6." if day == 1 else f"TASK 1 · Part 1\n{blk(blocks, '1.2')}.") + "\nFormat: Handout (or on the playground)"
    t[4] = f"COACHING\nTASK 1 · {blk(blocks, '1.3')}.\nTRAINER document, section Coaching.\nFormat: Trainer only + Handout"
    t[5] = (
        f"CASE STUDY {case} · 2 UE\nTASK 1 · Part 2: {blk(blocks, '2.1')}; {blk(blocks, '2.2')}.\n"
        "UE5: brief and own work. UE6: groups compare; the model solution is shown at the end (TRAINER document, section Case study).\n"
        "Format: Handout; model solution shown after the work"
    )
    t[7] = "FEEDBACK ROUND\nTRAINER document, section Feedback round (groups present their reasoning; Level 3 input questions).\nFormat: Presentation by the groups; Trainer only"
    t[8] = (
        f"TRANSFER PROJECT · {role} · 3 UE\nUE8: MATERI B, Core cards {ids(b_core)}. Format: Presentation + Handout.\n"
        f"UE9: TASK 2, {blk(blocks, '3.1')}. Format: Handout.\n"
        f"UE10: TASK 2, {blk(blocks, '3.2')}; then debrief with the model solution (TRAINER document, section Transfer project). Format: Handout; model solution shown after the work"
    )
    return t


def lines(text, width):
    return sum(max(1, math.ceil(len(x) / width)) for x in text.split("\n"))


def build(out):
    wb = openpyxl.load_workbook(PLAN)
    en = wb["English"]
    if SHEET in wb.sheetnames:
        del wb[SHEET]
    ws = wb.create_sheet(SHEET)
    ws.sheet_properties.tabColor = "1750A8"
    st = lambda ref: en[ref]
    S_HEAD_L, S_HEAD = st("A1"), st("D1")
    S_LAB_R, S_TOPIC = st("A2"), st("D2")
    S_UE, S_LVL, S_CELL = st("A4"), st("B4"), st("D4")

    def put(cell, value, src, wrap=True):
        cell.value = value
        cell._style = copy(src._style)
        cell.alignment = Alignment(horizontal=src.alignment.horizontal, vertical="top", wrap_text=wrap)

    # header
    put(ws["A1"], "UE", S_HEAD_L)
    put(ws["B1"], "Level", S_HEAD_L)
    put(ws["C1"], "Methodology", S_HEAD_L)
    for d in range(1, 9):
        put(ws.cell(1, 3 + d), f"Documents - Day {d}", S_HEAD)

    # row 2: topic of the day (copied from the plan), row 3: documents
    ws.merge_cells("A2:C2")
    put(ws["A2"], "Topic of the Day", S_LAB_R)
    for c in ("B2", "C2"):
        ws[c]._style = copy(S_LAB_R._style)
    ws.merge_cells("A3:C3")
    put(ws["A3"], "Documents (version 1.1)", S_LAB_R)
    for c in ("B3", "C3"):
        ws[c]._style = copy(S_LAB_R._style)
    for d in range(1, 9):
        put(ws.cell(2, 3 + d), en.cell(2, 3 + d).value, S_TOPIC)
        m, t, tr = FILES[d]
        put(ws.cell(3, 3 + d), f"Materi: Day{d}-Materi-{P}-{m}.docx\nTask: Day{d}-Task-{P}-{t}.docx\nTrainer (not for participants): Day{d}-Trainer-{P}-{tr}.docx", S_CELL)

    # rows 4 to 13: UE1 to UE10, the same rows and merges as the plan
    labels = {r: (en.cell(r, 1).value, en.cell(r, 2).value, en.cell(r, 3).value) for r in range(4, 14)}
    for r in range(4, 14):
        ue, lvl, meth = labels[r]
        put(ws.cell(r, 1), ue, S_UE)
        put(ws.cell(r, 2), lvl, S_LVL)
        put(ws.cell(r, 3), meth, S_LVL)
        for d in range(1, 9):
            ws.cell(r, 3 + d)._style = copy(S_CELL._style)
    for rng in ("B8:B9", "C8:C9", "B11:B13", "C11:C13"):
        ws.merge_cells(rng)
    ws["B8"].value = "Level 2 Application (Case study, 2 UE)"
    ws["B11"].value = "Senior Level 3 Transfer Project (3 UE)"
    for d in range(1, 9):
        t = texts(d)
        col = 3 + d
        for ue, row in ((1, 4), (2, 5), (3, 6), (4, 7), (5, 8), (7, 10), (8, 11)):
            put(ws.cell(row, col), t[ue], S_CELL)
        ws.merge_cells(start_row=8, start_column=col, end_row=9, end_column=col)
        ws.merge_cells(start_row=11, start_column=col, end_row=13, end_column=col)
        ws.cell(9, col)._style = copy(S_CELL._style)
        for r in (12, 13):
            ws.cell(r, col)._style = copy(S_CELL._style)

    # row 14: playground addresses, row 15: approval
    ws.merge_cells("A14:C14")
    put(ws["A14"], "Playground (website)", S_LAB_R)
    for c in ("B14", "C14"):
        ws[c]._style = copy(S_LAB_R._style)
    ws.merge_cells("A15:C15")
    put(ws["A15"], "Version, date, approval", S_LAB_R)
    for c in ("B15", "C15"):
        ws[c]._style = copy(S_LAB_R._style)
    for d in range(1, 9):
        put(ws.cell(14, 3 + d), release.day_url(d), S_CELL)
        put(ws.cell(15, 3 + d), f"Version {release.VERSION} · {release.RELEASE_DATE}\nReview and approval: {release.REVIEWER}, {release.REVIEWER_ROLE}", S_CELL)

    # legend
    legend = [
        ("How to read this sheet", None),
        ("MATERI", "The study material of the day. Materi A serves Route 1 (Levels 1 and 2), Materi B serves Route 2 (Level 3). Participants get it as a handout; the trainer presents the Core cards on the playground."),
        ("TASK", "The worksheet of the day. Task 1 is Route 1 (Part 1 = Level 1, Part 2 = the Level 2 case study, 2 UE). Task 2 is Route 2, the Level 3 transfer project (3 UE). It carries no answer keys and no scores. Participants get it as a handout or work in it on the playground."),
        ("TRAINER", "Coaching, feedback round, model solutions (Musterlösung), assessment criteria (Bewertungskriterien) and the transition from Level 2 to Level 3. Not handed out to participants. Model solutions are shown only after the participants have worked."),
        ("Presentation / Handout", "Presentation = what the trainer shows on screen (the cards' scan line, figures and key boxes; the case; the model solution after the work). Handout = what participants read, fill in and keep."),
        ("PLAYGROUND", f"The course website ({release.SITE}): in class it is the place to practise and learn. Each day has its own address: {release.SITE}day/N/."),
        ("UE", f"One teaching unit (Unterrichtseinheit) is counted as {release.UE_MINUTES} minutes; breaks are not shown. The case study is 2 UE (UE5–6), the transfer project 3 UE (UE8–10)."),
    ]
    r = 17
    for k, v in legend:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        put(ws.cell(r, 1), k, S_UE if v is None else S_LVL)
        for c in (2, 3):
            ws.cell(r, c)._style = copy((S_UE if v is None else S_LVL)._style)
        if v is not None:
            ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=11)
            put(ws.cell(r, 4), v, S_CELL)
            for c in range(5, 12):
                ws.cell(r, c)._style = copy(S_CELL._style)
            ws.row_dimensions[r].height = 15 * max(2, math.ceil(len(v) / 190)) + 4
        else:
            ws.row_dimensions[r].height = 24
        r += 1

    # widths and heights
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 30
    for d in range(1, 9):
        ws.column_dimensions[openpyxl.utils.get_column_letter(3 + d)].width = 46
    ws.row_dimensions[1].height = 30
    for row in (2, 3, 14, 15):
        h = max(lines(str(ws.cell(row, 3 + d).value or ""), 48) for d in range(1, 9))
        ws.row_dimensions[row].height = 15 * h + 6
    for row in (4, 5, 6, 7, 10):
        h = max(lines(str(ws.cell(row, 3 + d).value or ""), 48) for d in range(1, 9))
        ws.row_dimensions[row].height = 15 * h + 6
    for row, span in ((8, 2), (11, 3)):
        h = max(lines(str(ws.cell(row, 3 + d).value or ""), 48) for d in range(1, 9))
        for rr in range(row, row + span):
            ws.row_dimensions[rr].height = (15 * h + 8) / span
    ws.freeze_panes = "D2"
    ws.sheet_view.zoomScale = 85
    wb.save(out)
    return out


if __name__ == "__main__":
    print(build(sys.argv[1]))
