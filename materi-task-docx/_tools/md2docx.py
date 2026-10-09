"""
Markdown dialect -> .docx for the AION Digital Learning Materi and Task documents.
Dialect (see DOCX-EXPORT-GUIDE.md section 4):
  front matter  --- key: value ... ---
  # / ## / ### / ####  headings          **bold** *italic* `code`
  - bullet     1. numbered     > scan line     ^src: sources line
  | pipe | tables |  (<br> = line break in a cell, a cell that is exactly __ is a blank writing space)
  ![caption](fig/name.png)
  :::box Title ... :::     mist box ("In plain words")
  :::rules Title ... :::   decision-rules box (accent)
  :::note Title ... :::    outlined box
  :::warn Title ... :::    rust box (case assumption)
  [[ANSWER lines=n]]  [[OPTIONS]] a || b   [[CHOOSE]] a || b   [[TICKALL]] a || b
  [[CONTENTS]]  ---pagebreak---   %% reviewer notes (not printed)
"""
import os
import re
import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
from glossary import GLOSSARY  # noqa: E402
from refs import REFS  # noqa: E402

INK = "17212E"
ASH = "556274"
ACC = "1750A8"
ACCSOFT = "E3ECFA"
MIST = "E6ECF4"
LINE = "D5DEE9"
SIG = "0B6F69"
SIGSOFT = "DCF0EE"
RUST = "AD3F26"
RUSTSOFT = "F8E4DE"
CANVAS = "F3F6FA"
FONT = "Calibri"

TEXT_W = 9638  # dxa, A4 with 2 cm margins

ORDER = {
    "pPr": "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange".split(),
    "rPr": "rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof snapToGrid vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd fitText vertAlign rtl cs em lang eastAsianLayout specVanish oMath".split(),
    "tcPr": "cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark".split(),
    "tblPr": "tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc tblCellSpacing tblInd tblBorders shd tblLayout tblCellMar tblLook tblCaption tblDescription".split(),
    "trPr": "cnfStyle divId gridBefore gridAfter wBefore wAfter cantSplit trHeight tblHeader tblCellSpacing jc hidden".split(),
    "tcBorders": "top start left bottom end right insideH insideV tl2br tr2bl".split(),
    "tblBorders": "top start left bottom end right insideH insideV".split(),
    "tcMar": "top start left bottom end right".split(),
    "tblCellMar": "top start left bottom end right".split(),
}


def fix_order(root):
    for tag, order in ORDER.items():
        for el in root.iter(qn("w:" + tag)):
            kids = list(el)
            key = lambda k: order.index(k.tag.split("}")[1]) if k.tag.split("}")[1] in order else 999
            srt = sorted(kids, key=key)
            if srt != kids:
                for k in kids:
                    el.remove(k)
                for k in srt:
                    el.append(k)


def el(tag, **attrs):
    e = OxmlElement("w:" + tag)
    for k, v in attrs.items():
        e.set(qn("w:" + k), str(v))
    return e


def shade(cell_or_par_pr, fill):
    s = el("shd", val="clear", color="auto", fill=fill)
    cell_or_par_pr.append(s)


# ------------------------------------------------------------------ inline

INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`|<br>)")


def add_inline(par, text, size=None, color=None, bold=False, italic=False):
    parts = INLINE.split(text)
    for p in parts:
        if not p:
            continue
        if p == "<br>":
            par.add_run().add_break(WD_BREAK.LINE)
            continue
        b, i, mono = bold, italic, False
        if p.startswith("**") and p.endswith("**") and len(p) > 4:
            p, b = p[2:-2], True
        elif p.startswith("*") and p.endswith("*") and len(p) > 2:
            p, i = p[1:-1], True
        elif p.startswith("`") and p.endswith("`") and len(p) > 2:
            p, mono = p[1:-1], True
        r = par.add_run(p)
        r.bold = b or None
        r.italic = i or None
        if mono:
            r.font.name = "Consolas"
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
    return par


def par_fmt(par, before=0, after=4, line=None, keep_next=False, keep_lines=False, align=None, left=None, hanging=None):
    pf = par.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line:
        pf.line_spacing = line
    if keep_next:
        pf.keep_with_next = True
    if keep_lines:
        pf.keep_together = True
    if align == "center":
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if left is not None:
        pf.left_indent = Cm(left)
    if hanging is not None:
        pf.first_line_indent = Cm(-hanging)


# ------------------------------------------------------------------ parsing

def parse_front(text):
    meta = {}
    m = re.match(r"^---\n(.*?)\n---\n", text, re.S)
    if not m:
        return meta, text
    for ln in m.group(1).split("\n"):
        if ":" in ln:
            k, v = ln.split(":", 1)
            meta[k.strip()] = v.strip()
    return meta, text[m.end():]


def parse_blocks(lines):
    """lines -> list of blocks (tuples)"""
    blocks = []
    i = 0
    n = len(lines)
    while i < n:
        ln = lines[i]
        s = ln.strip()
        if not s or s.startswith("%%"):
            i += 1
            continue
        if s == "---pagebreak---":
            blocks.append(("pagebreak",))
            i += 1
            continue
        if s == "[[CONTENTS]]":
            blocks.append(("contents",))
            i += 1
            continue
        m = re.match(r"^:::(box|rules|note|warn)\s*(.*)$", s)
        if m:
            kind, title = m.group(1), m.group(2)
            depth = 1
            j = i + 1
            inner = []
            while j < n:
                t = lines[j].strip()
                if re.match(r"^:::(box|rules|note|warn)\b", t):
                    depth += 1
                elif t == ":::":
                    depth -= 1
                    if depth == 0:
                        break
                inner.append(lines[j])
                j += 1
            blocks.append(("box", kind, title, parse_blocks(inner)))
            i = j + 1
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            blocks.append(("h", len(m.group(1)), m.group(2)))
            i += 1
            continue
        m = re.match(r"^!\[(.*?)\]\((.*?)\)$", s)
        if m:
            blocks.append(("img", m.group(1), m.group(2)))
            i += 1
            continue
        m = re.match(r"^\[\[ANSWER(?:\s+lines=(\d+))?(?:\s+label=(.*?))?\]\]$", s)
        if m:
            blocks.append(("answer", int(m.group(1) or 3), m.group(2)))
            i += 1
            continue
        m = re.match(r"^\[\[(OPTIONS|CHOOSE|TICKALL)\]\]\s*(.*)$", s)
        if m:
            blocks.append(("choice", m.group(1), [x.strip() for x in re.split(r"\|\||‖", m.group(2))]))
            i += 1
            continue
        if s.startswith("^src:"):
            blocks.append(("src", s[5:].strip()))
            i += 1
            continue
        if s.startswith(">"):
            txt = []
            while i < n and lines[i].strip().startswith(">"):
                txt.append(lines[i].strip().lstrip(">").strip())
                i += 1
            blocks.append(("scan", " ".join(txt)))
            continue
        if s.startswith("|"):
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip())
                i += 1
            cells = [[c.strip() for c in r.strip("|").split("|")] for r in rows]
            cells = [r for r in cells if not all(re.match(r"^:?-{2,}:?$", c) for c in r)]
            blocks.append(("table", cells))
            continue
        if re.match(r"^[-*] ", s):
            items = []
            while i < n and re.match(r"^\s*[-*] ", lines[i]):
                items.append(re.sub(r"^\s*[-*] ", "", lines[i]).strip())
                i += 1
            blocks.append(("bullets", items))
            continue
        if re.match(r"^\d+\. ", s):
            items = []
            while i < n and re.match(r"^\s*\d+\. ", lines[i]):
                items.append(re.sub(r"^\s*\d+\. ", "", lines[i]).strip())
                i += 1
            blocks.append(("numbered", items))
            continue
        # paragraph: join following non-special lines
        buf = [s]
        i += 1
        while i < n:
            t = lines[i].strip()
            if not t or t.startswith(("#", "|", ">", "![", ":::", "[[", "^src:", "%%", "---pagebreak", "**")) or re.match(r"^([-*]|\d+\.) ", t):
                break
            buf.append(t)
            i += 1
        blocks.append(("p", " ".join(buf)))
    return blocks


# ------------------------------------------------------------------ rendering

class Builder:
    def __init__(self, meta, base_dir):
        self.meta = meta
        self.base = base_dir
        self.doc = Document()
        self.fig_n = 0
        self.headings = []
        self._setup()

    # ---- setup
    def _setup(self):
        d = self.doc
        sec = d.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.0)
        sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
        st = d.styles
        n = st["Normal"]
        n.font.name = FONT
        n.font.size = Pt(10.5)
        n.font.color.rgb = RGBColor.from_string(INK)
        rpr = n.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = el("rFonts")
            rpr.append(rf)
        for a in ("ascii", "hAnsi", "eastAsia", "cs"):
            rf.set(qn("w:" + a), FONT)
        n.paragraph_format.space_after = Pt(4)
        n.paragraph_format.line_spacing = 1.12
        for name, size, color, before, after in (("Heading 1", 17, INK, 18, 6), ("Heading 2", 13.5, INK, 14, 4), ("Heading 3", 11.5, INK, 10, 3), ("Heading 4", 10.5, ASH, 8, 2)):
            h = st[name]
            h.font.name = FONT
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.italic = False
            h.font.color.rgb = RGBColor.from_string(color)
            r = h.element.get_or_add_rPr()
            rf = r.find(qn("w:rFonts"))
            if rf is None:
                rf = el("rFonts")
                r.append(rf)
            for a in ("ascii", "hAnsi", "eastAsia", "cs"):
                rf.set(qn("w:" + a), FONT)
            for a in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
                if rf.get(qn("w:" + a)):
                    del rf.attrib[qn("w:" + a)]
            h.paragraph_format.space_before = Pt(before)
            h.paragraph_format.space_after = Pt(after)
            h.paragraph_format.keep_with_next = True
        for name in ("List Bullet",):
            b = st[name]
            b.font.name = FONT
            b.font.size = Pt(10.5)
            b.paragraph_format.space_after = Pt(2)
        self._footer(sec)

    def _footer(self, sec):
        f = sec.footer
        p = f.paragraphs[0]
        p.text = ""
        txt = self.meta.get("footer") or f"Habmann AufstiegsAkademie · Digital Learning · Day {self.meta.get('day', '')} · {self.meta.get('doc', '').title()}"
        r = p.add_run(txt + "   ·   Page ")
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor.from_string(ASH)
        for code in ("PAGE",):
            self._field(p, code)
        r = p.add_run(" of ")
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor.from_string(ASH)
        self._field(p, "NUMPAGES")
        pPr = p._p.get_or_add_pPr()
        bd = el("pBdr")
        bd.append(el("top", val="single", sz="4", space="4", color=LINE))
        pPr.append(bd)

    def _field(self, p, code):
        r = p.add_run()
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor.from_string(ASH)
        b = el("fldChar", fldCharType="begin")
        i = el("instrText")
        i.set(qn("xml:space"), "preserve")
        i.text = f" {code} "
        s = el("fldChar", fldCharType="separate")
        t = el("t")
        t.text = "1"
        e = el("fldChar", fldCharType="end")
        for x in (b, i, s, t, e):
            r._r.append(x)

    # ---- low level table helpers
    def _tbl(self, container, rows, cols, widths):
        t = container.add_table(rows=rows, cols=cols)
        tblPr = t._tbl.tblPr
        for ch in list(tblPr):
            if ch.tag in (qn("w:tblStyle"), qn("w:tblW"), qn("w:tblLayout"), qn("w:tblLook")):
                tblPr.remove(ch)
        tblPr.append(el("tblW", w=sum(widths), type="dxa"))
        tblPr.append(el("tblLayout", type="fixed"))
        grid = t._tbl.tblGrid
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(w))
        for r in t.rows:
            for c, w in zip(r.cells, widths):
                tcPr = c._tc.get_or_add_tcPr()
                tcw = tcPr.find(qn("w:tcW"))
                if tcw is None:
                    tcw = el("tcW")
                    tcPr.append(tcw)
                tcw.set(qn("w:w"), str(w))
                tcw.set(qn("w:type"), "dxa")
        return t

    def _borders(self, tblPr, color=LINE, sz=4, inner=True):
        b = el("tblBorders")
        for side in ("top", "left", "bottom", "right"):
            b.append(el(side, val="single", sz=sz, space=0, color=color))
        if inner:
            b.append(el("insideH", val="single", sz=sz, space=0, color=color))
            b.append(el("insideV", val="single", sz=sz, space=0, color=color))
        tblPr.append(b)

    def _cellmar(self, tblPr, top=60, bottom=60, left=100, right=100):
        m = el("tblCellMar")
        for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
            m.append(el(side, w=v, type="dxa"))
        tblPr.append(m)

    def _cell_shade(self, cell, fill):
        tcPr = cell._tc.get_or_add_tcPr()
        tcPr.append(el("shd", val="clear", color="auto", fill=fill))

    def _cell_borders(self, cell, **sides):
        tcPr = cell._tc.get_or_add_tcPr()
        b = el("tcBorders")
        for side, (color, sz) in sides.items():
            b.append(el(side, val="single", sz=sz, space=0, color=color))
        tcPr.append(b)

    def _row_height(self, row, cm, rule="atLeast"):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(el("trHeight", val=int(cm * 567), hRule=rule))

    def _cant_split(self, row):
        row._tr.get_or_add_trPr().append(el("cantSplit"))

    def _first_par(self, cell):
        return cell.paragraphs[0]

    # ---- blocks
    def render(self, blocks, container, width=TEXT_W):
        for b in blocks:
            getattr(self, "b_" + b[0])(b, container, width)

    def _add_par(self, container, first_cell_par=False):
        if first_cell_par and hasattr(container, "paragraphs") and len(container.paragraphs) == 1 and not container.paragraphs[0].text and not getattr(container, "_used", False):
            container._used = True
            return container.paragraphs[0]
        return container.add_paragraph()

    def b_h(self, b, c, w):
        _, lvl, text = b
        p = self._add_par(c, True) if hasattr(c, "_tc") else c.add_paragraph()
        p.style = self.doc.styles[f"Heading {min(lvl, 4)}"]
        add_inline(p, text)
        if lvl == 1:
            pPr = p._p.get_or_add_pPr()
            bd = el("pBdr")
            bd.append(el("bottom", val="single", sz="8", space="2", color=INK))
            pPr.append(bd)
        if lvl <= 2 and not hasattr(c, "_tc"):
            self.headings.append((lvl, re.sub(r"\*", "", text)))

    def b_p(self, b, c, w):
        p = self._add_par(c, True)
        add_inline(p, b[1])
        par_fmt(p, after=5)

    def b_scan(self, b, c, w):
        p = self._add_par(c, True)
        add_inline(p, b[1], size=11, color=ASH, italic=True)
        par_fmt(p, before=2, after=6, left=0.3)
        pPr = p._p.get_or_add_pPr()
        bd = el("pBdr")
        bd.append(el("left", val="single", sz="18", space="6", color=ACC))
        pPr.append(bd)

    def b_src(self, b, c, w):
        p = self._add_par(c, True)
        r = p.add_run("Sources · ")
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor.from_string(ASH)
        add_inline(p, b[1], size=9, color=ASH)
        par_fmt(p, before=4, after=8)

    def b_bullets(self, b, c, w):
        for it in b[1]:
            p = self._add_par(c, True)
            p.style = self.doc.styles["List Bullet"]
            add_inline(p, it)
            par_fmt(p, after=2)
        # small gap after list
        c.paragraphs[-1].paragraph_format.space_after = Pt(5)

    def b_numbered(self, b, c, w):
        for k, it in enumerate(b[1], 1):
            p = self._add_par(c, True)
            add_inline(p, f"{k}.\t" + it)
            par_fmt(p, after=2, left=0.75, hanging=0.75)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(0.75))
        c.paragraphs[-1].paragraph_format.space_after = Pt(5)

    def b_pagebreak(self, b, c, w):
        c.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def b_contents(self, b, c, w):
        self.contents_marker = c.add_paragraph()  # filled after the body is rendered

    def b_img(self, b, c, w):
        _, cap, path = b
        full = os.path.join(self.base, path)
        if not os.path.exists(full):
            full = os.path.join(ROOT, path)
        if not os.path.exists(full):
            raise FileNotFoundError(full)
        self.fig_n += 1
        from PIL import Image

        with Image.open(full) as im:
            wpx, hpx = im.size
        max_w = min(16.5, (w / 567) - 0.4)
        width_cm = max_w
        height_cm = width_cm * hpx / wpx
        if height_cm > 15.5:
            height_cm = 15.5
            width_cm = height_cm * wpx / hpx
        p = self._add_par(c, True)
        par_fmt(p, before=4, after=2, keep_next=True, align="center")
        p.add_run().add_picture(full, width=Cm(width_cm))
        cp = c.add_paragraph()
        add_inline(cp, f"Figure {self.fig_n} · {cap}", size=9, color=ASH, italic=True)
        par_fmt(cp, after=8, align="center")

    def b_answer(self, b, c, w):
        _, lines, label = b
        if label:
            lp = self._add_par(c, True)
            add_inline(lp, label, size=9.5, color=ASH, bold=True)
            par_fmt(lp, before=2, after=1, keep_next=True)
        t = self._tbl(c, 1, 1, [w - 40])
        self._borders(t._tbl.tblPr, "8793A3", 6, False)
        self._cellmar(t._tbl.tblPr)
        self._row_height(t.rows[0], 0.62 * lines + 0.3)
        self._cant_split(t.rows[0])
        gap = c.add_paragraph()
        par_fmt(gap, after=4)
        for r in gap.runs:
            r.font.size = Pt(4)

    def b_choice(self, b, c, w):
        _, kind, items = b
        head = {"OPTIONS": "Tick one", "CHOOSE": "Tick one", "TICKALL": "Tick all that apply"}[kind]
        p = self._add_par(c, True)
        r = p.add_run(head + ":  ")
        r.italic = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor.from_string(ASH)
        if kind == "OPTIONS":
            for k, it in enumerate(items):
                add_inline(p, "☐ " + it + ("      " if k < len(items) - 1 else ""))
            par_fmt(p, after=6)
        else:
            par_fmt(p, after=1, keep_next=True)
            for it in items:
                q = c.add_paragraph()
                add_inline(q, "☐ " + it)
                par_fmt(q, after=1, left=0.6, hanging=0.6)
            c.paragraphs[-1].paragraph_format.space_after = Pt(6)

    def b_table(self, b, c, w):
        rows = b[1]
        ncol = max(len(r) for r in rows)
        rows = [r + [""] * (ncol - len(r)) for r in rows]
        # column widths: every column first gets room for its longest word, the rest is shared by average text length
        sz = 9.5 if ncol > 3 else 10
        cd = 104 if sz < 10 else 108
        mins, avgs = [], []
        for j in range(ncol):
            col = [re.sub(r"\*", "", r[j]) for r in rows]
            words = [x for cell in col for part in re.split(r"<br>", cell) for x in part.split()]
            lw = max((len(x) for x in words), default=1)
            mins.append(lw * cd + 240)
            lines = [max((len(p) for p in re.split(r"<br>", cell)), default=0) for cell in col[1:]] or [0]
            avgs.append(max(sum(lines) / len(lines), 4))
            if all(r[j].strip() in ("__", "") for r in rows[1:]):
                mins[-1] = max(mins[-1], 1500)
                avgs[-1] = max(avgs[-1], 14)
        if sum(mins) > w:
            k = w / sum(mins)
            mins = [m * k for m in mins]
        extra = w - sum(mins)
        tot = sum(avgs)
        widths = [int(m + extra * a / tot) for m, a in zip(mins, avgs)]
        widths[-1] += w - sum(widths)
        t = self._tbl(c, len(rows), ncol, widths)
        self._borders(t._tbl.tblPr)
        self._cellmar(t._tbl.tblPr)
        for i, r in enumerate(rows):
            self._cant_split(t.rows[i])
            if i == 0:
                t.rows[i]._tr.get_or_add_trPr().append(el("tblHeader"))
            for j, txt in enumerate(r):
                cell = t.cell(i, j)
                p = cell.paragraphs[0]
                if txt.strip() == "__":
                    self._row_height(t.rows[i], 0.95)
                    continue
                add_inline(p, txt, size=9.5 if ncol > 3 else 10, bold=(i == 0))
                par_fmt(p, after=0)
                if i == 0:
                    self._cell_shade(cell, MIST)
        gap = c.add_paragraph()
        par_fmt(gap, after=4)

    def b_box(self, b, c, w):
        _, kind, title, inner = b
        fill, bar = {"box": (MIST, ASH), "rules": (ACCSOFT, ACC), "note": ("FFFFFF", "8793A3"), "warn": (RUSTSOFT, RUST)}[kind]
        if kind == "box" and title.lower().startswith("how to decide"):
            kind, fill, bar = "rules", ACCSOFT, ACC
        t = self._tbl(c, 1, 1, [w - 40])
        pr = t._tbl.tblPr
        self._borders(pr, LINE if kind != "note" else "8793A3", 4, False)
        self._cellmar(pr, 90, 90, 160, 160)
        cell = t.cell(0, 0)
        self._cell_shade(cell, fill)
        self._cell_borders(cell, left=(bar, 28))
        cell._used = False
        if title:
            p = cell.paragraphs[0]
            cell._used = True
            r = p.add_run(title.upper() if kind in ("box", "warn") else title)
            r.bold = True
            r.font.size = Pt(9 if kind in ("box", "warn") else 10)
            r.font.color.rgb = RGBColor.from_string(bar if kind != "box" else ASH)
            par_fmt(p, after=3, keep_next=True)
        self.render(inner, cell, w - 360)
        gap = c.add_paragraph()
        par_fmt(gap, after=4)

    # ---- document parts
    def cover(self):
        m = self.meta
        d = self.doc
        p = d.add_paragraph()
        add_inline(p, "HABMANN AUFSTIEGSAKADEMIE · DIGITAL LEARNING · EDUCATIONAL UX/UI DESIGN", size=8.5, color=ASH, bold=True)
        par_fmt(p, after=14)
        p = d.add_paragraph()
        add_inline(p, m.get("kicker", ""), size=10, color=ACC, bold=True)
        par_fmt(p, after=2)
        p = d.add_paragraph()
        add_inline(p, m.get("title", ""), size=26, bold=True)
        par_fmt(p, after=4, line=1.0)
        if m.get("subtitle"):
            p = d.add_paragraph()
            add_inline(p, m["subtitle"], size=13, color=ASH)
            par_fmt(p, after=10)
        rows = [
            ("Day", m.get("daytitle") or f"Day {m.get('day')}"),
            ("Document", m.get("doctype", "")),
            ("Route and level", m.get("route", "")),
            ("Time", m.get("minutes", "")),
            ("Website (playground)", m.get("website", "")),
            ("Case", m.get("case", "")),
        ]
        rows = [r for r in rows if r[1]]
        t = self._tbl(d, len(rows), 2, [2300, TEXT_W - 2300])
        self._borders(t._tbl.tblPr)
        self._cellmar(t._tbl.tblPr, 70, 70, 120, 120)
        for i, (k, v) in enumerate(rows):
            a, bcell = t.cell(i, 0), t.cell(i, 1)
            self._cell_shade(a, MIST)
            add_inline(a.paragraphs[0], k, size=9.5, bold=True, color=ASH)
            add_inline(bcell.paragraphs[0], v, size=10)
            par_fmt(a.paragraphs[0], after=0)
            par_fmt(bcell.paragraphs[0], after=0)
        d.add_paragraph().paragraph_format.space_after = Pt(4)
        if m.get("intro"):
            self.render([("box", "box", "About this document", [("p", m["intro"])])], d)

    def contents(self):
        if not getattr(self, "contents_marker", None):
            return
        marker = self.contents_marker
        items = [h for h in self.headings if h[0] == 2]
        parent = marker._p.getparent()
        idx = list(parent).index(marker._p)
        # build a small box table at the marker position
        tmp = self.doc.add_paragraph()  # temp anchor to build content in the body, then move
        t = self._tbl(self.doc, 1, 1, [TEXT_W - 40])
        self._borders(t._tbl.tblPr, LINE, 4, False)
        self._cellmar(t._tbl.tblPr, 90, 90, 160, 160)
        cell = t.cell(0, 0)
        self._cell_shade(cell, CANVAS)
        self._cell_borders(cell, left=(ASH, 28))
        p = cell.paragraphs[0]
        r = p.add_run("IN THIS DOCUMENT")
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor.from_string(ASH)
        par_fmt(p, after=3)
        for lvl, txt in items:
            q = cell.add_paragraph()
            add_inline(q, txt, size=10)
            par_fmt(q, after=1, left=0.3)
        parent.remove(t._tbl)
        parent.insert(idx + 1, t._tbl)
        parent.remove(tmp._p)

    def glossary_appendix(self):
        keys = [k.strip() for k in self.meta.get("glossary", "").split(",") if k.strip()]
        if not keys:
            return
        d = self.doc
        miss = [k for k in keys if k not in GLOSSARY]
        if miss:
            raise KeyError(f"glossary keys missing: {miss}")
        d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        self.render([("h", 1, "Glossary")], d)
        self.render([("p", "Every term below is used in this document. Terms stay plain text in the body; this list says what they mean in everyday words.")], d)
        rows = [["Term", "In plain words", "Example", "Comes from"]]
        for k in sorted(keys, key=lambda x: GLOSSARY[x]["term"].lower()):
            g = GLOSSARY[k]
            rows.append([f"**{g['term']}**", g["plain"], g.get("example", "—"), g.get("from", "—")])
        self.render([("table", rows)], d)

    def references_appendix(self):
        keys = [k.strip() for k in self.meta.get("refs", "").split(",") if k.strip()]
        if not keys:
            return
        miss = [k for k in keys if k not in REFS]
        if miss:
            raise KeyError(f"ref keys missing: {miss}")
        d = self.doc
        self.render([("h", 1, "References")], d)
        self.render([("p", "Check each source and its current edition before you teach from it. Standards and laws marked as in flux can change.")], d)
        for k in keys:
            p = d.add_paragraph()
            add_inline(p, REFS[k], size=9.5)
            par_fmt(p, after=4, left=0.5, hanging=0.5)

    def save(self, path):
        fix_order(self.doc.element)
        for sec in self.doc.sections:
            fix_order(sec.footer._element)
        # styles
        fix_order(self.doc.styles.element)
        z = self.doc.settings.element.find(qn("w:zoom"))
        if z is not None and z.get(qn("w:percent")) is None:
            z.set(qn("w:percent"), "100")
        self.doc.core_properties.title = self.meta.get("title", "")
        self.doc.core_properties.author = "AION Digital Learning"
        self.doc.core_properties.subject = self.meta.get("kicker", "")
        self.doc.save(path)


def build(md_path, out_path):
    text = open(md_path, encoding="utf-8").read()
    meta, body = parse_front(text)
    base = os.path.dirname(md_path)
    b = Builder(meta, ROOT)
    b.cover()
    blocks = parse_blocks(body.split("\n"))
    b.render(blocks, b.doc)
    b.references_appendix()
    b.glossary_appendix()
    b.contents()
    b.save(out_path)
    return out_path


if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2])
    print("built", sys.argv[2])
